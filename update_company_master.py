"""Refresh the searchable TSE listed-company master from JPX's official Excel list."""
import io
import json
import re
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

JPX_URL = 'https://www.jpx.co.jp/markets/statistics-equities/misc/tvdivq0000001vg2-att/data_j.xlsx'
OUT = Path('data/company_master.json')


def clean(v):
    return str(v or '').strip()


def find_header(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True):
        vals = [clean(v) for v in row]
        if 'コード' in vals and '銘柄名' in vals and '市場・商品区分' in vals:
            return vals.index('コード'), vals.index('銘柄名'), vals.index('市場・商品区分')
    raise RuntimeError('JPX Excel header not found')


def keep_company(market):
    m = clean(market)
    if not m:
        return False
    excluded = ('ETF', 'ETN', 'REIT', 'リート', 'インフラファンド', '優先株式', '投資法人')
    if any(x in m.upper() for x in excluded):
        return False
    # Ordinary listed companies, including domestic/foreign stocks and TOKYO PRO Market.
    return any(x in m for x in ('プライム', 'スタンダード', 'グロース', 'PRO Market'))


def normalize_code(v):
    s = clean(v).upper().replace('.0', '')
    if re.fullmatch(r'[0-9A-Z]{4,5}', s):
        return s
    if re.fullmatch(r'\d{1,4}', s):
        return s.zfill(4)
    return None


def main():
    req = urllib.request.Request(JPX_URL, headers={'User-Agent': 'Mozilla/5.0 kabu-score'})
    with urllib.request.urlopen(req, timeout=60) as r:
        raw = r.read()
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    code_i, name_i, market_i = find_header(ws)

    companies = []
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = normalize_code(row[code_i] if code_i < len(row) else '')
        name = clean(row[name_i] if name_i < len(row) else '')
        market = clean(row[market_i] if market_i < len(row) else '')
        if not code or not name or code in seen or not keep_company(market):
            continue
        seen.add(code)
        companies.append({'code': code, 'name': name, 'market': market})

    companies.sort(key=lambda x: x['code'])
    if len(companies) < 3000:
        raise RuntimeError(f'JPX master unexpectedly small: {len(companies)}')

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({'source': 'JPX', 'source_url': JPX_URL, 'companies': companies}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Updated company master: {len(companies)} companies')


if __name__ == '__main__':
    main()
